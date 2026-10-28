"""Run the preserved final workflow offline, with isolated output and source evidence."""
import hashlib
import importlib
import json
from pathlib import Path
import socket
import sys
import tempfile
from decimal import Decimal


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inventory(root):
    result = {}
    for folder in ['每次付款后修改_持续更新数据库','月度修改_月度账单','输出指引']:
        for path in sorted((root/folder).rglob('*')):
            if path.is_file() and path.suffix.lower() in ('.xlsx','.xls','.json','.xmind','.png','.jpg','.jpeg','.pdf','.webp','.bmp') and not path.name.startswith('~$'):
                result[str(path.relative_to(root))] = sha(path)
    return result


def event(label):
    print(json.dumps({'event':'progress','label':label}, ensure_ascii=False), flush=True)


def main(request_path):
    request_file = Path(request_path).resolve()
    job = request_file.parent
    request = json.loads(request_file.read_text(encoding='utf-8'))
    source = Path(request['source_root'])
    # The historical workflow is local-only. No model, OCR download or outbound sockets.
    def deny(*args, **kwargs):
        raise RuntimeError('尾期工作流禁止网络/API访问')
    socket.socket.connect = deny
    socket.socket.connect_ex = deny
    socket.create_connection = deny
    sys.path.insert(0, request['bundle'])
    receipt = dict(status='blocked',workflow_id=request['workflow_id'],job_id=request['job_id'],
                   building_id=request['building_id'],source_kind='原Excel工作流资料；不是数据库月度金额重算',
                   real_api_calls=0,key_read=False,output_is_review_draft=True)
    target = job/'statement.xlsx'
    try:
        event('核对原工作流资料指纹')
        before = inventory(source)
        receipt['source_files_sha256'] = before
        final = importlib.import_module('对账单最终工作流')
        rules = importlib.import_module('水电板块工作流')
        rent = importlib.import_module('租金Excel接入')
        args = rules.build_parser().parse_args(['generate','--root',str(source),'--building-id',request['building_id']])
        # Resolve identity once before any output; original workflow independently rechecks it.
        app_path, _, _ = rent.discover(source)
        book = rent.SourceBook(app_path)
        try: app = rent.resolve_application(book,args,importlib.import_module('租金板块工作流'))
        finally: book.close()
        receipt['application'] = {k:str(app[k]) for k in ['building_id','project','building','exit_date','exit_type','application_rows','tail_paid','remarks']}
        receipt['application_source'] = dict(file=str(app_path.relative_to(source)),sheet=app['application_sheet'],rows=app['application_rows'])
        original_rent = rent.generate
        def rental(*args, **kwargs):
            try: return original_rent(*args, **kwargs)
            except Exception:
                report = Path(args[0].output).with_suffix('.诊断.json')
                if report.is_file():
                    receipt['diagnostic'] = json.loads(report.read_text(encoding='utf-8'))
                raise
        rent.generate = rental
        with tempfile.TemporaryDirectory(prefix='tail_',dir=job) as temp:
            args.output = str(Path(temp)/'statement.xlsx')
            original_water = final.water.generate
            original_finish = final.summary.finish
            def water(*a, **kw):
                event('执行租金与水电工作流（仅本地OCR，缺失会标注）')
                return original_water(*a, **kw)
            def finish(*a, **kw):
                event('汇总费用、高风险提示及取数记录')
                return original_finish(*a, **kw)
            final.water.generate = water
            final.summary.finish = finish
            result = final.generate(args)
            event('校验生成文件及来源未发生变化')
            if inventory(source) != before:
                raise ValueError('生成过程中来源资料已改变，请重新核对版本后再生成。')
            from openpyxl import load_workbook
            produced = Path(result['output'])
            wb = load_workbook(produced,read_only=True,data_only=True)
            try:
                final_amount = Decimal(str(result['final_building_share']))
                fee = Decimal(str(result['management_fee']))
                if not final_amount.is_finite() or not fee.is_finite() or final_amount*100 != (final_amount*100).to_integral_value():
                    raise ValueError('工作流金额非法或超过分精度。')
                cached_final = Decimal(str(wb.worksheets[0]['H6'].value))
                cached_fee = Decimal(str(wb.worksheets[0]['G6'].value))
                if cached_final != final_amount or cached_fee != fee:
                    raise ValueError('工作流金额与保存后工作簿缓存不一致。')
                if '取数记录' not in wb.sheetnames:
                    raise ValueError('生成结果缺少取数记录。')
                cell = wb.worksheets[0]
                if rent.name_key(cell['A6'].value) != rent.name_key(app['project']) or rent.name_key(cell['B6'].value) != rent.name_key(app['building']):
                    raise ValueError('生成账单对象与已识别楼栋不一致。')
            finally: wb.close()
            produced.replace(target)
            result.pop('output',None)
        receipt.update(status='completed',workflow_version=final.WORKFLOW_VERSION,result=result,output_sha256=sha(target),sources_unchanged=True)
    except Exception as exc:
        target.unlink(missing_ok=True)
        # Known data issues contain the actual source/sheet/row evidence, not credentials.
        receipt['error'] = '\n'.join(line for line in str(exc).splitlines() if not line.startswith('诊断文件：'))[:3000] if isinstance(exc,(ValueError,OSError)) else '处理异常：'+type(exc).__name__
    (job/'receipt.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding='utf-8')
    return 0 if receipt['status']=='completed' else 2


if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
