"""R35只读核查：不修改业务库、来源文件或正式入口，不读取Key。"""
from collections import Counter
import contextlib
from datetime import datetime,timezone,timedelta
from decimal import Decimal
import json
from pathlib import Path
import sys

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from agent.settings import Settings,ROOT
from agent.r35_database import R35Database
from agent.database import sha256
from openpyxl import load_workbook


def verify():
    s=Settings.load();db=R35Database(s);binding=db.verify_binding();directory=s.database.parent
    pointer_paths=['数据表/当前版本.json','数据表/当前正式Excel版本.json','数据库结构/当前结构版本.json']
    pointers={p:sha256(s.project_root/p) for p in pointer_paths}
    manifest=json.loads((directory/'R35正式发布文件清单.json').read_text(encoding='utf8'))
    hashes={name:sha256(directory/name)==digest for name,digest in manifest.items()}
    with contextlib.closing(db.connect()) as c:
        integrity=c.execute('PRAGMA integrity_check').fetchone()[0];fk=c.execute('PRAGMA foreign_key_check').fetchall()
        monthly=[dict(r) for r in c.execute('SELECT * FROM T08_月度金额分')]
        successful=[r[0] for r in c.execute('SELECT "付款金额（元）" FROM "T08-2" WHERE "付款状态（备注）"=\'支付成功\'')]
        formula=all(r['我方应付分']==r['我方实收分']+r['调整净额分']-r['应收业主费用分']-r['水电预留分']-r['水费划扣分']-r['电费划扣分'] for r in monthly if r['我方应付分'] is not None)
        columns=[r[1] for r in c.execute('PRAGMA table_info(T08)')]
        db_rows=[tuple(r) for r in c.execute('SELECT * FROM T08')]
    book=directory/'T08_楼栋月度我方应付_R35_20261017.xlsx';wb=load_workbook(book,read_only=True,data_only=True)
    sheet=next((ws for ws in wb if list(next(ws.iter_rows(min_row=1,max_row=1,values_only=True)))==columns),None)
    if sheet is None:raise RuntimeError('未找到与T08字段一致的工作表')
    excel_rows=[tuple(r) for r in sheet.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
    amount_fields={'租赁实收','赔偿实收','水费实收','电费实收','我方实收合计','租赁调整','赔偿调整','应收费用调整','水电预留费用调整','计入调整合计','暂不计入水电调整原额','应收业主费用','水电预留费用','水费划扣','电费划扣','我方应付金额'}
    def normalize(row):
        return tuple(None if v is None or v=='' else format(Decimal(str(v)).normalize(),'f') if columns[i] in amount_fields else str(v) for i,v in enumerate(row))
    workbook_equal=Counter(map(normalize,db_rows))==Counter(map(normalize,excel_rows));wb.close()
    totals={key:sum(r[key] or 0 for r in monthly) for key in ['我方实收分','应收业主费用分','我方应付分']}
    expected={'我方实收分':13595928543,'应收业主费用分':1278826058,'我方应付分':11691136905}
    unchanged=all(sha256(s.project_root/p)==digest for p,digest in pointers.items())
    result={'checked_at':datetime.now(timezone(timedelta(hours=8))).isoformat(),'binding':binding,'fingerprints':hashes,'pointer_hashes':pointers,'pointers_unchanged':unchanged,
            'integrity_check':integrity,'foreign_key_check':fk,'monthly_rows':len(monthly),'unknown_payable_rows':sum(r['我方应付分'] is None for r in monthly),
            'totals_cents':totals,'formula_all_computed_rows_equal':formula,'excel_rows':len(excel_rows),'excel_matches_database':workbook_equal,
            'successful_documents':len(successful),'successful_amount_yuan':str(sum(map(Decimal,successful),Decimal(0))),'business_database_written':False,'real_api_calls':0}
    result['passed']=all(hashes.values()) and unchanged and integrity=='ok' and not fk and formula and workbook_equal and totals==expected and len(monthly)==9576 and len(successful)==3860
    (ROOT/'runtime').mkdir(exist_ok=True);(ROOT/'runtime/r35_agent_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in result.items() if k not in ['fingerprints','pointer_hashes','binding']},ensure_ascii=False,indent=2))
    return result


if __name__=='__main__':sys.exit(0 if verify()['passed'] else 1)
