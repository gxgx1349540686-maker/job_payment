"""只读验收R33。数据目录内零写入，不自动导入、不更新正式指针。"""
import contextlib
import json
import pathlib
import sqlite3
import sys
from collections import Counter
from decimal import Decimal
from datetime import datetime, timezone

sys.path.insert(0,str(pathlib.Path(__file__).resolve().parents[1]))
from agent.settings import ROOT, Settings
from agent.database import sha256
from openpyxl import load_workbook


def verify():
    s=Settings.load(); directory=s.database.parent
    manifest=json.loads((directory/'交付文件清单.json').read_text(encoding='utf-8-sig'))
    fingerprints={n:sha256(directory/n)==h for n,h in manifest.items()}
    validation=json.loads(s.validation.read_text(encoding='utf-8-sig'))
    expected_sources={
        'R31数据库':s.project_root/'数据表/尾期结算迁移_R31_20261017/job_payment_OBS_V16_R31_20261017.sqlite',
        'R31工作簿':s.project_root/'数据表/尾期结算迁移_R31_20261017/尾期结算数据表_R31_20261017.xlsx',
        'R32数据库':s.project_root/'数据表/尾期水电与退出_R32_20261017/job_payment_OBS_V17_R32_20261017.sqlite'}
    fingerprints.update({k:sha256(p)==validation['source_hashes'][k] for k,p in expected_sources.items()})
    source_manifest=json.loads((s.project_root/'数据表/尾期结算迁移_R31_20261017/来源及去重清单.json').read_text(encoding='utf-8-sig'))
    for source in source_manifest['sources']:
        fingerprints[source['file']]=sha256(pathlib.Path(source['file']))==source['sha256']
    pointers={}
    for path in ['数据表/当前版本.json','数据表/当前正式Excel版本.json','数据库结构/当前结构版本.json']:
        pointer=s.project_root/path; data=json.loads(pointer.read_text(encoding='utf-8-sig'))
        pointers[path]={'version':data.get('version'),'sha256':sha256(pointer)}
    formal=json.loads((s.project_root/'数据表/当前版本.json').read_text(encoding='utf-8-sig'))
    fingerprints['正式R30数据库']=sha256(pathlib.Path(formal['database_file']))==formal['sha256']
    results={}
    with contextlib.closing(sqlite3.connect(s.database.as_uri()+'?mode=ro',uri=True)) as c:
        integrity=c.execute('PRAGMA integrity_check').fetchall()
        fk=c.execute('PRAGMA foreign_key_check').fetchall()
        for table,info in validation['workbooks'].items():
            columns=[row[1] for row in c.execute(f'PRAGMA table_info("{table}")')]
            db_rows=c.execute(f'SELECT * FROM "{table}"').fetchall()
            wb=load_workbook(directory/info['file'],read_only=True,data_only=False)
            main=None
            for ws in wb:
                head=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
                if list(head)==columns:
                    main=ws;break
            if main is None:raise RuntimeError(f'{table}找不到同字段主表')
            excel_rows=[tuple(row) for row in main.iter_rows(min_row=2,values_only=True) if any(v is not None for v in row)]
            def normalize(row):
                return tuple(None if v is None or v=='' else (format(Decimal(str(v)).normalize(),'f') if i==3 and table=='T06-4' else str(v)) for i,v in enumerate(row))
            same=Counter(map(normalize,excel_rows))==Counter(map(normalize,db_rows))
            results[table]={'database_rows':len(db_rows),'excel_rows':len(excel_rows),'excel_sheet':main.title,'all_business_values_equal':same}
            wb.close()
        sums={}
        for category,value in c.execute('SELECT 科目,金额 FROM "T06-4"'):
            sums[category]=sums.get(category,Decimal(0))+Decimal(value)
        counts={'T07':c.execute('select count(*) from T07').fetchone()[0],
                'successful_payments':c.execute('select count(*) from "T08-2" where "付款状态（备注）"=\'支付成功\'').fetchone()[0],
                'exit_matched':c.execute('select count(*) from "T05-3" where "新模式合同id" is not null and "新模式合同id"<>\'\'').fetchone()[0],
                'source_rows':c.execute('select count(*) from "尾期衍生__来源记录对照"').fetchone()[0]}
        payments=sum((Decimal(r[0]) for r in c.execute('select "付款金额（元）" from "T08-2" where "付款状态（备注）"=\'支付成功\'')),Decimal(0))
    unchanged_pointers=all(sha256(s.project_root/p)==d['sha256'] for p,d in pointers.items())
    business_ok=(counts=={'T07':5837,'successful_payments':3860,'exit_matched':369,'source_rows':916}
                 and payments==Decimal('112529007.69') and sums=={'水费':Decimal('274253.39'),'电费':Decimal('1201545.00')})
    report={'checked_at':datetime.now(timezone.utc).isoformat(),'database':str(s.database),'fingerprints':fingerprints,
            'pointers':pointers,'pointers_unchanged':unchanged_pointers,'workbooks':results,'integrity_check':integrity,
            'foreign_key_check':fk,'counts':counts,'water_totals':{k:str(v) for k,v in sums.items()},'payment_total':str(payments),
            'database_write_performed':False,'import_needed':not all(x['all_business_values_equal'] for x in results.values()),
            'passed':all(fingerprints.values()) and integrity==[('ok',)] and not fk and unchanged_pointers and business_ok and all(x['all_business_values_equal'] for x in results.values())}
    target=ROOT/'runtime/r33_readonly_verification.json';target.parent.mkdir(exist_ok=True)
    target.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps({k:v for k,v in report.items() if k not in ['fingerprints','pointers']},ensure_ascii=False,indent=2))
    return report


if __name__=='__main__':
    sys.exit(0 if verify()['passed'] else 1)
