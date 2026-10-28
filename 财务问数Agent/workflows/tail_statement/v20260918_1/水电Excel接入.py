"""水电Excel接入层。

处理顺序：先调用租金工作流生成中间结果，再读取普通月度水电、尾期水电、
付款实付、划扣填写值和凭证图片，最后把全部内容写入新版模板。
"""
from __future__ import annotations

import argparse
import io
import json
import re
import tempfile
import zipfile
from collections import defaultdict, deque
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

import 租金Excel接入 as rent_excel
import 租金板块工作流 as rent_rules


NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
ROOT = Path(__file__).resolve().parent
TAIL_FOLDER = "每次付款后修改_持续更新数据库/需更新_楼栋尾期账单水电费取数表"
WATER_NAMES = ["水费", "水费收入", "水费总额"]
ELECTRIC_NAMES = ["电费", "电费收入", "电费总额"]
TAIL_MONTH_NAMES = ["收入归属月"]
TAIL_BEGIN_NAMES = ["费用开始时间", "合同开始时间", "合同开始日期"]
TAIL_END_NAMES = ["费用结束时间", "合同结束时间", "合同结束日期"]


class WaterIssue(ValueError):
    """水电取数或输出阶段的明确异常。"""


class ReserveUnavailable(WaterIssue):
    """源预留缺失可暂按0；重复记录等冲突仍须阻断。"""


def optional_index(headers, names):
    """字段可缺失，但同一含义出现多列时拒绝静默选取。"""
    wanted = {rent_excel.clean(name) for name in names}
    hits = [i for i, value in enumerate(headers) if rent_excel.clean(value) in wanted]
    if len(hits) > 1:
        raise WaterIssue(f"字段重复：{names}，匹配{len(hits)}列")
    return hits[0] if hits else None


def value_amount(value, formula, context, rules):
    """普通空白表示该项未发生；无缓存公式不能当0。"""
    if value is None or (isinstance(value, str) and not value.strip()):
        if formula:
            raise WaterIssue(f"{context}：公式没有计算缓存，请在Excel中重算保存")
        return Decimal("0.00")
    try:
        return rules.amount(value)
    except ValueError as exc:
        raise WaterIssue(f"{context}：{exc}") from exc


def read_application(root, args):
    """复用租金关联键，确定唯一申请行和退出月份。"""
    app_path, history_path, month_paths = rent_excel.discover(root)
    book = rent_excel.SourceBook(app_path)
    try:
        app = rent_excel.resolve_application(book, args, rent_rules)
    finally:
        book.close()
    return app_path, history_path, month_paths, app


def zero_water_note(rows, month, source):
    """零收入只备注，不阻断；无记录与有记录合计为零分别说明。"""
    if not rows:
        return f"{month}月{source}未取得非零水电明细，水电收入按0计"
    if all(sum((row[key] for row in rows), Decimal(0)) == 0 for key in ('water', 'electric')):
        return f"{month}月{source}水费、电费收入合计均为0"
    return ""


def read_regular_month_water(path, bill_month, app, rules):
    """读取N月普通明细的拆分水电，并归属为N-1月收入。"""
    if bill_month <= 1:
        raise WaterIssue("普通水电收入从2月账单开始读取，并归属为前一月")
    income_month = bill_month - 1
    book = rent_excel.SourceBook(path)
    try:
        sheet = rent_excel.DETAIL_NAMES.get(bill_month, "明细表")
        headers, rows = rent_excel.table(book, sheet, ["id", "room", "lease", "comp"])
        id_i = rent_excel.index(headers, rent_excel.ALIASES["id"])
        room_i = rent_excel.index(headers, rent_excel.ALIASES["room"])
        begin_i = optional_index(headers, rent_excel.ALIASES["begin"])
        end_i = optional_index(headers, rent_excel.ALIASES["end"])
        water_i = optional_index(headers, WATER_NAMES)
        electric_i = optional_index(headers, ELECTRIC_NAMES)
        if water_i is None or electric_i is None:
            raise WaterIssue(f"{path.name}/{sheet}：未同时找到拆分的水费和电费字段")
        output = []
        for row_no, row, formulas in rows:
            if not rent_excel.same_id(rent_excel.cell(row, id_i), app["building_id"]):
                continue
            water = value_amount(
                rent_excel.cell(row, water_i), formulas.get(water_i, ""),
                f"{path.name}/{sheet}!{get_column_letter(water_i + 1)}{row_no}", rules,
            )
            electric = value_amount(
                rent_excel.cell(row, electric_i), formulas.get(electric_i, ""),
                f"{path.name}/{sheet}!{get_column_letter(electric_i + 1)}{row_no}", rules,
            )
            if water == 0 and electric == 0:
                continue
            output.append({
                "project": app["project"], "building": app["building"],
                "room": str(rent_excel.cell(row, room_i) or ""),
                "begin": rent_excel.cell(row, begin_i), "end": rent_excel.cell(row, end_i),
                "lease": Decimal("0.00"), "comp": Decimal("0.00"),
                "water": water, "electric": electric,
                "remark": "", "month": f"{app['exit_date'].year}-{income_month:02d}",
                "content": f"{app['exit_date'].year}-{income_month:02d}水电（取自{bill_month}月普通明细）",
                "file": str(path), "sheet": sheet, "row": row_no,
                "source_kind": "普通月度水电", "source_bill_month": bill_month,
            })
        return output, sheet
    finally:
        book.close()


def read_source_reserve(path, month, app, rules):
    """读取月度汇总中的源水电预留，1月兼容“水电扣留费用”。"""
    book = rent_excel.SourceBook(path)
    try:
        sheet = rent_excel.SUMMARY_NAMES.get(month, "汇总表")
        headers, rows = rent_excel.table(book, sheet, ["id", "lease", "comp"])
        id_i = rent_excel.index(headers, rent_excel.ALIASES["id"])
        hits = [(n, row, formulas) for n, row, formulas in rows
                if rent_excel.same_id(rent_excel.cell(row, id_i), app["building_id"])]
        if not hits:
            raise ReserveUnavailable(f'{path.name}/{sheet}：未取得该楼栋源水电预留记录')
        if len(hits) != 1:
            raise WaterIssue(f"{path.name}/{sheet}：楼栋汇总匹配{len(hits)}条")
        row_no, row, formulas = hits[0]
        names = ["水电扣留费用", "水电预留费用"] if month == 1 else ["水电预留费用"]
        reserve_i = optional_index(headers, names)
        if reserve_i is None:
            raise ReserveUnavailable(f"{path.name}/{sheet}：缺少水电预留字段")
        reserve = value_amount(
            rent_excel.cell(row, reserve_i), formulas.get(reserve_i, ""),
            f"{path.name}/{sheet}!{get_column_letter(reserve_i + 1)}{row_no}", rules,
        )
        return reserve, sheet, f"{get_column_letter(reserve_i + 1)}{row_no}", formulas.get(reserve_i, "")
    finally:
        book.close()


def read_monthly_paid(history_path, app, exit_month):
    """付款申请实付是水电预留是否计入的唯一开关。"""
    book = rent_excel.SourceBook(history_path)
    try:
        sheet = rent_rules.select_history_sheet(list(book.sheets))
        headers, rows = rent_excel.table(book, sheet, ["id"])
        id_i = rent_excel.index(headers, rent_excel.ALIASES["id"])
        hits = [(n, row, formulas) for n, row, formulas in rows
                if rent_excel.same_id(rent_excel.cell(row, id_i), app["building_id"])]
        if len(hits) != 1:
            raise WaterIssue(f"{history_path.name}/{sheet}：楼栋匹配{len(hits)}条")
        row_no, row, formulas = hits[0]
        result = {}
        for month in range(1, exit_month + 1):
            try:
                payment = rent_rules.history_payment(headers, row, month, formulas)
            except ValueError as exc:
                raise WaterIssue(f"{month}月实付读取失败：{exc}") from exc
            result[month] = payment
        return result, sheet, row_no
    finally:
        book.close()


def filename_date(path, default_year):
    """按文件名日期排序：YYYYMMDD、YYMMDD、MMDD，无法识别返回None。"""
    name = path.stem
    for raw in re.findall(r"(?<!\d)(20\d{6})(?!\d)", name):
        try:
            return datetime.strptime(raw, "%Y%m%d").date()
        except ValueError:
            pass
    for raw in re.findall(r"(?<!\d)(\d{6})(?!\d)", name):
        try:
            return datetime.strptime("20" + raw, "%Y%m%d").date()
        except ValueError:
            pass
    for raw in re.findall(r"(?<!\d)(\d{4})(?!\d)", name):
        try:
            return date(default_year, int(raw[:2]), int(raw[2:]))
        except ValueError:
            pass
    return None


def tail_sort_key(path, default_year):
    parsed = filename_date(path, default_year)
    if parsed is not None:
        return 1, parsed.toordinal(), path.stat().st_mtime
    return 0, 0, path.stat().st_mtime


def month_matches(value, target):
    try:
        parsed = rent_excel.as_date(value)
        return parsed.year == target.year and parsed.month == target.month
    except rent_excel.DataIssue:
        text = str(value or "")
        match = re.search(r"(20\d{2})[-/年.]?(\d{1,2})", text)
        return bool(match and int(match.group(1)) == target.year and int(match.group(2)) == target.month)


def scan_tail_file(path, app, rules):
    """扫描一个尾期文件，返回该楼栋全部记录及X月记录。"""
    book = rent_excel.SourceBook(path)
    building_rows = []
    target_rows = []
    try:
        for sheet in book.sheets:
            rows = book.rows(sheet)
            for j, (header_row, headers, _) in enumerate(rows[:20]):
                room_i = optional_index(headers, rent_excel.ALIASES["room"])
                month_i = optional_index(headers, TAIL_MONTH_NAMES)
                water_i = optional_index(headers, WATER_NAMES)
                electric_i = optional_index(headers, ELECTRIC_NAMES)
                if None in (room_i, month_i) or (water_i is None and electric_i is None):
                    continue
                id_i = optional_index(headers, rent_excel.ALIASES["id"])
                project_i = optional_index(headers, rent_excel.ALIASES["project"])
                building_i = optional_index(headers, rent_excel.ALIASES["building"])
                begin_i = optional_index(headers, TAIL_BEGIN_NAMES)
                end_i = optional_index(headers, TAIL_END_NAMES)
                for row_no, row, formulas in rows[j + 1:]:
                    id_hit = id_i is not None and rent_excel.same_id(rent_excel.cell(row, id_i), app["building_id"])
                    name_hit = (
                        building_i is not None
                        and (id_i is None or not rent_excel.normalize_id(rent_excel.cell(row, id_i)))
                        and rent_excel.name_key(rent_excel.cell(row, building_i)) in [rent_excel.name_key(n) for n in app.get('building_aliases', [app['building']])]
                        and (project_i is None or rent_excel.name_key(rent_excel.cell(row, project_i)) == rent_excel.name_key(app["project"]))
                    )
                    if not (id_hit or name_hit):
                        continue
                    building_rows.append((sheet, row_no))
                    if not month_matches(rent_excel.cell(row, month_i), app["exit_date"]):
                        continue
                    water = value_amount(
                        rent_excel.cell(row, water_i), formulas.get(water_i, "") if water_i is not None else "",
                        f"{path.name}/{sheet}!{get_column_letter((water_i or 0) + 1)}{row_no}", rules,
                    ) if water_i is not None else Decimal("0.00")
                    electric = value_amount(
                        rent_excel.cell(row, electric_i), formulas.get(electric_i, "") if electric_i is not None else "",
                        f"{path.name}/{sheet}!{get_column_letter((electric_i or 0) + 1)}{row_no}", rules,
                    ) if electric_i is not None else Decimal("0.00")
                    if water == 0 and electric == 0:
                        continue
                    target_rows.append({
                        "project": app["project"], "building": app["building"],
                        "room": str(rent_excel.cell(row, room_i) or ""),
                        "begin": rent_excel.cell(row, begin_i), "end": rent_excel.cell(row, end_i),
                        "lease": Decimal("0.00"), "comp": Decimal("0.00"),
                        "water": water, "electric": electric, "remark": "",
                        "month": f"{app['exit_date']:%Y-%m}",
                        "content": f"{app['exit_date']:%Y-%m}尾期水电",
                        "file": str(path), "sheet": sheet, "row": row_no,
                        "source_kind": "尾期水电",
                    })
                break
        return building_rows, target_rows
    finally:
        book.close()


def read_tail_water(root, app, rules, extra_path=None):
    """先找含该楼栋的文件，再按文件名日期选择最新版本。"""
    folder = root / Path(TAIL_FOLDER)
    candidates = []
    paths = list(folder.glob('*.xlsx')) if folder.exists() else []
    if extra_path: paths.append(Path(extra_path).resolve())
    if paths:
        for path in dict.fromkeys(paths):
            if path.name.startswith("~$"):
                continue
            building_rows, target_rows = scan_tail_file(path, app, rules)
            if building_rows:
                candidates.append((path, building_rows, target_rows))
    if not candidates:
        return [], None, "未取得尾期水电"
    selected = max(candidates, key=lambda item: tail_sort_key(item[0], app["exit_date"].year))
    path, _, rows = selected
    if not rows:
        return [], path, f"未取得{app['exit_date'].month}月尾期水电"
    return rows, path, ""


def read_application_actuals_and_images(app_path, app, rules):
    """读取申请表1—4月填写金额及同一申请行的嵌入图片。"""
    workbook = load_workbook(app_path, data_only=True, keep_links=False)
    try:
        sheet = workbook[app["application_sheet"]]
        headers = [cell.value for cell in sheet[1]]
        row_no = app["application_rows"][-1]
        water = {}
        electric = {}
        source_cells = {}
        for month in range(1, 5):
            for kind, label, target in (
                ("水费", f"{month}月划扣水费", water),
                ("电费", f"{month}月划扣电费", electric),
            ):
                indexes = [i for i, value in enumerate(headers) if rent_excel.clean(value) == rent_excel.clean(label)]
                if len(indexes) != 1:
                    raise WaterIssue(f"{app_path.name}/{app['application_sheet']}：字段{label}匹配{len(indexes)}列")
                index = indexes[0]
                raw = sheet.cell(row_no, index + 1).value
                target[month] = Decimal("0.00") if raw in (None, "") else rules.amount(raw)
                source_cells[(kind, month)] = f"{get_column_letter(index + 1)}{row_no}"
        images = []
        for image in sheet._images:
            try:
                anchor_row = image.anchor._from.row + 1
            except AttributeError:
                continue
            if anchor_row not in app["application_rows"]:
                continue
            data = image._data()
            images.append({
                "data": data,
                "format": (image.format or "png").lower(),
                "width": image.width,
                "height": image.height,
                "anchor_row": anchor_row,
            })
        return water, electric, source_cells, images
    finally:
        workbook.close()


def ocr_images(images, skip=False):
    """本地OCR；依赖不可用时不阻止出账，但明确标注未自动核验。"""
    if skip or not images:
        return [], "已跳过OCR" if skip else "没有申请行凭证图片"
    try:
        import easyocr
        import numpy as np
        from PIL import Image

        reader = easyocr.Reader(["ch_sim", "en"], gpu=False, download_enabled=False, verbose=False)
        results = []
        for number, item in enumerate(images, 1):
            if item.get('text'):
                text = item['text']
            else:
                source = np.array(Image.open(io.BytesIO(item["data"])).convert("RGB"))
                pieces = reader.readtext(source, detail=0, paragraph=False)
                text = "\n".join(str(piece) for piece in pieces)
            water_score = text.count("水费") + text.count("用水量") * 2 + text.count("污水")
            electric_score = text.count("电费") * 2 + text.count("用电")
            kind = "水费" if water_score > electric_score else "电费"
            results.append({"number": number, "kind": kind, "text": text})
        return results, ""
    except Exception as exc:
        return [], f"OCR未完成：{type(exc).__name__}: {exc}"


def parse_money_tokens(text):
    tokens = []
    for raw in re.findall(r"(?<!\d)(\d{1,3}(?:,\d{3})*(?:\.\d{1,2})|\d+\.\d{1,2})(?!\d)", text):
        try:
            tokens.append(Decimal(raw.replace(",", "")))
        except Exception:
            continue
    return tokens


def parse_water_sms(ocr_results, exit_year, rules):
    """短信凭证按扣费月份的前一个月归属水费；无法结构化的文字仍留在审计页。"""
    voucher = {}
    for result in ocr_results:
        if result["kind"] != "水费":
            continue
        text = re.sub(r"\s+", "", result["text"])
        pattern = re.compile(
            r"合计([\d,]+(?:\.\d{1,2})?)元[\s\S]{0,100}?于(20\d{2})年(\d{1,2})月(\d{1,2})日"
        )
        for amount_text, year_text, month_text, _ in pattern.findall(text):
            paid_year, paid_month = int(year_text), int(month_text)
            bill_month = 12 if paid_month == 1 else paid_month - 1
            bill_year = paid_year - 1 if paid_month == 1 else paid_year
            if bill_year != exit_year:
                continue
            voucher[bill_month] = rules.amount(amount_text.replace(",", ""))
    return voucher


def compare_actuals(water, electric, ocr_results, exit_month, rules):
    """填写金额是主金额；OCR只负责一致性核验和5月后凭证补充。"""
    notes = []
    water_voucher = parse_water_sms(ocr_results, 2026 if exit_month is None else exit_month.year, rules)
    text_by_kind = defaultdict(str)
    for result in ocr_results:
        text_by_kind[result["kind"]] += "\n" + result["text"]
    token_by_kind = {kind: parse_money_tokens(text) for kind, text in text_by_kind.items()}
    limit = exit_month.month
    actual_water = {}
    actual_electric = {}
    for month in range(1, limit + 1):
        if month <= 4:
            actual_water[month] = water.get(month, Decimal("0.00"))
            actual_electric[month] = electric.get(month, Decimal("0.00"))
        else:
            actual_water[month] = water_voucher.get(month, Decimal("0.00"))
            actual_electric[month] = Decimal("0.00")

        if month <= 4:
            filled = actual_water[month]
            voucher = water_voucher.get(month)
            if voucher is None:
                if filled != 0:
                    notes.append(f"{month}月水费填写{filled}，OCR未取得对应凭证金额")
            elif abs(filled - voucher) > rules.OCR_TOLERANCE:
                notes.append(f"{month}月水费填写{filled}与凭证{voucher}不一致，按填写金额计入")

            filled_electric = actual_electric[month]
            if filled_electric != 0:
                matches = token_by_kind.get("电费", [])
                if not any(abs(filled_electric - value) <= rules.OCR_TOLERANCE for value in matches):
                    notes.append(f"{month}月电费填写{filled_electric}，OCR未取得对应凭证金额")
    return actual_water, actual_electric, water_voucher, notes


def normalize_date(value):
    if value in (None, ""):
        return None
    try:
        return rent_excel.as_date(value)
    except rent_excel.DataIssue:
        return str(value)


def detail_key(row):
    return (
        str(row.get("month") or ""), str(row.get("room") or ""),
        normalize_date(row.get("begin")), normalize_date(row.get("end")),
    )


def read_rent_output(path):
    """把租金中间结果转换为统一明细，不重算租金业务。"""
    workbook = load_workbook(path, data_only=False, keep_links=False)
    values_workbook = load_workbook(path, data_only=True, keep_links=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        values_sheet = values_workbook[values_workbook.sheetnames[0]]
        details = []
        for row_no in range(10, sheet.max_row + 1):
            project = sheet.cell(row_no, 1).value
            building = sheet.cell(row_no, 2).value
            if not project or not building:
                continue
            content = str(sheet.cell(row_no, 11).value or "")
            month_match = re.match(r"(20\d{2}-\d{2})", content)
            details.append({
                "project": project, "building": building,
                "room": str(sheet.cell(row_no, 3).value or ""),
                "begin": sheet.cell(row_no, 4).value, "end": sheet.cell(row_no, 5).value,
                "lease": Decimal(str(sheet.cell(row_no, 6).value or 0)),
                "comp": Decimal(str(sheet.cell(row_no, 7).value or 0)),
                "water": Decimal("0.00"), "electric": Decimal("0.00"),
                "remark": sheet.cell(row_no, 8).value or "",
                "month": month_match.group(1) if month_match else "",
                "content": content, "source_kind": "租金结果",
            })
        rent_audit = []
        if "取数记录" in workbook.sheetnames:
            audit = workbook["取数记录"]
            rent_audit = [[cell.value for cell in row] for row in audit.iter_rows()]
        return {
            "details": details,
            "room_count": values_sheet["C6"].value or sheet["C6"].value or 0,
            "lease_total": Decimal(str(values_sheet["D6"].value or sum((d["lease"] for d in details), Decimal(0)))),
            "comp_total": Decimal(str(values_sheet["E6"].value or sum((d["comp"] for d in details), Decimal(0)))),
            "audit": rent_audit,
            "notes": [str(row[2]) for row in rent_audit if len(row)>2 and row[0]=='租金缺失备注'],
        }
    finally:
        workbook.close()
        values_workbook.close()


def merge_details(rent_details, water_details):
    """同月、同房间、同起止日期优先合并；其余作为水电专用行追加。"""
    output = [dict(row) for row in rent_details]
    indexes = defaultdict(deque)
    for number, row in enumerate(output):
        indexes[detail_key(row)].append(number)
    for water_row in water_details:
        key = detail_key(water_row)
        if indexes[key]:
            target = output[indexes[key].popleft()]
            target["water"] += water_row["water"]
            target["electric"] += water_row["electric"]
            if "水电" not in target["content"]:
                target["content"] += "＋水电"
        else:
            output.append(dict(water_row))
    return output


def copy_cell(source, target):
    target.value = source.value
    if source.has_style:
        target._style = copy(source._style)
    target.number_format = source.number_format
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def capture_bottom(sheet, start=28, end=38):
    cells = []
    for row in range(start, end + 1):
        for col in range(1, 11):
            source = sheet.cell(row, col)
            cells.append((row - start, col, source.value, copy(source._style), source.number_format,
                          copy(source.font), copy(source.fill), copy(source.border),
                          copy(source.alignment), copy(source.protection)))
    merges = [str(item) for item in sheet.merged_cells.ranges
              if item.min_row >= start and item.max_row <= end and item.max_col <= 10]
    heights = {row - start: sheet.row_dimensions[row].height for row in range(start, end + 1)}
    return cells, merges, heights


def restore_bottom(sheet, captured, destination):
    cells, merges, heights = captured
    for relative_row, col, value, style, number_format, font, fill, border, alignment, protection in cells:
        target = sheet.cell(destination + relative_row, col)
        target.value = value
        target._style = copy(style)
        target.number_format = number_format
        target.font = copy(font)
        target.fill = copy(fill)
        target.border = copy(border)
        target.alignment = copy(alignment)
        target.protection = copy(protection)
    for relative_row, height in heights.items():
        sheet.row_dimensions[destination + relative_row].height = height
    for address in merges:
        match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", address)
        if not match:
            continue
        col1, row1, col2, row2 = match.groups()
        shift = destination - 28
        sheet.merge_cells(f"{col1}{int(row1) + shift}:{col2}{int(row2) + shift}")


def patch_formula_cache(path, values):
    """openpyxl不计算公式；为预览软件写入已独立核算的缓存。"""
    with zipfile.ZipFile(path) as archive:
        contents = {name: archive.read(name) for name in archive.namelist()}
    xml = ET.fromstring(contents["xl/worksheets/sheet1.xml"])
    for cell in xml.iter(NS + "c"):
        address = cell.attrib.get("r")
        if address not in values:
            continue
        value = cell.find(NS + "v")
        if value is None:
            value = ET.SubElement(cell, NS + "v")
        value.text = str(values[address])
    contents["xl/worksheets/sheet1.xml"] = ET.tostring(xml, encoding="utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in contents.items():
            archive.writestr(name, data)


def add_images(sheet, images, ocr_results):
    """凭证按水费/电费分两栏保留原始图片；返回图片覆盖到的末行。"""
    result_kind = {item["number"]: item["kind"] for item in ocr_results}
    for address in list(sheet.merged_cells.ranges):
        if address.min_row >= 16 and address.min_col >= 13:
            sheet.unmerge_cells(str(address))
    for row in range(16, 80):
        for col in range(13, 30):
            sheet.cell(row, col).value = None
            sheet.cell(row, col).fill = PatternFill()
            sheet.cell(row, col).border = Border()
    sheet["M16"] = "水费划扣凭证（申请表/补充凭证）"
    sheet["U16"] = "电费划扣凭证（申请表/补充凭证）"
    for coordinate in ("M16", "U16"):
        sheet[coordinate].font = Font(bold=True, color="7F6000")
        sheet[coordinate].fill = PatternFill("solid", fgColor="FFF2CC")
    next_row = {"水费": 17, "电费": 17}
    buffers = []
    for number, item in enumerate(images, 1):
        kind = result_kind.get(number, "电费")
        buffer = io.BytesIO(item["data"])
        buffers.append(buffer)
        image = XLImage(buffer)
        max_width = 420 if kind == "水费" else 620
        max_height = 1000 if kind == "水费" else 520
        scale = min(max_width / item["width"], max_height / item["height"], 1)
        image.width = int(item["width"] * scale)
        image.height = int(item["height"] * scale)
        column = "M" if kind == "水费" else "U"
        anchor = f"{column}{next_row[kind]}"
        sheet.add_image(image, anchor)
        next_row[kind] += max(2, int(image.height / 48) + 2)
    sheet._water_image_buffers = buffers
    return max(next_row.values())


def write_output(template, output, app, rent_data, details, monthly_values, notes,
                 images, ocr_results, trace, rules):
    """保留新版模板总账单和水电结算区，扩展房间明细与凭证区域。"""
    if output.exists():
        raise WaterIssue(f"输出已存在，请更换--output路径：{output}")
    workbook = load_workbook(template, keep_links=False)
    chosen = "xx(项目名)_xx(楼栋名)"
    if chosen not in workbook.sheetnames:
        raise WaterIssue(f"模板缺少工作表：{chosen}")
    for name in list(workbook.sheetnames):
        if name != chosen:
            del workbook[name]
    sheet = workbook[chosen]
    sheet.title = re.sub(r"[\\/*?:\[\]]", "_", app["project"] + "_" + app["building"])[:31]

    bottom = capture_bottom(sheet)
    for address in list(sheet.merged_cells.ranges):
        if address.min_row >= 28 and address.max_col <= 10:
            sheet.unmerge_cells(str(address))

    base_styles = [copy(sheet.cell(10, col)._style) for col in range(1, 12)]
    for row in range(10, max(39, 9 + len(details)) + 1):
        for col in range(1, 12):
            cell = sheet.cell(row, col)
            cell.value = None
            cell._style = copy(base_styles[col - 1])

    sheet["A1"] = f"{app['project']} {app['building']} 对账单"
    sheet["A2"] = f"对账周期：2026年{app['exit_date'].month}月"
    sheet["A3"] = f"所属门店及楼栋：{app['project']} {app['building']}"
    sheet["A4"] = "一、楼栋总账单"
    sheet["A6"], sheet["B6"] = app["project"], app["building"]
    sheet["C5"] = f"{app['exit_date'].year}年{app['exit_date'].month}月\n在租房间数"
    sheet["C6"] = rent_data["room_count"]
    last_detail = max(10, 9 + len(details))
    sheet["D6"] = f"=SUM(F10:F{last_detail})"
    sheet["E6"] = f"=SUM(G10:G{last_detail})"
    sheet["F6"] = "=-N14"
    sheet["G6"] = 0
    sheet["H6"] = "=D6+E6-F6-G6"
    sheet["K6"] = "租金及水电部分已生成"
    sheet["K6"].alignment = Alignment(wrap_text=True, vertical="center")

    for row_no, item in enumerate(details, 10):
        values = [
            app["project"], app["building"], str(item.get("room") or ""),
            item.get("begin"), item.get("end"), float(item["lease"]), float(item["comp"]),
            item.get("remark") or "", float(item["water"]), float(item["electric"]), item.get("content") or "",
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_no, col)
            cell.value = value
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for col in (4, 5):
            try:
                sheet.cell(row_no, col).value = rent_excel.as_date(values[col - 1])
                sheet.cell(row_no, col).number_format = "yyyy-mm-dd"
            except rent_excel.DataIssue:
                pass
        for col in (6, 7, 9, 10):
            sheet.cell(row_no, col).number_format = "#,##0.00;[Red]-#,##0.00"
        sheet.row_dimensions[row_no].height = 38

    month_columns = {month: 14 + (month - 1) * 2 for month in range(1, 9)}
    for month, col in month_columns.items():
        values = monthly_values.get(month, {})
        sheet.cell(9, col).value = float(values.get("reserve", Decimal("0.00")))
        sheet.cell(10, col).value = float(values.get("electric_actual", Decimal("0.00")))
        sheet.cell(11, col).value = float(values.get("water_actual", Decimal("0.00")))
        sheet.cell(12, col).value = float(values.get("electric_income", Decimal("0.00")))
        sheet.cell(13, col).value = float(values.get("water_income", Decimal("0.00")))
        for row in range(9, 14):
            sheet.cell(row, col).number_format = "#,##0.00;[Red]-#,##0.00"
    for row in range(9, 14):
        cells = "+".join(f"{get_column_letter(col)}{row}" for col in month_columns.values())
        sheet[f"AC{row}"] = f"={cells}"
        sheet[f"AC{row}"].number_format = "#,##0.00;[Red]-#,##0.00"
    sheet["N14"] = "=AC9-(AC10+AC11)+(AC12+AC13)"
    sheet["N14"].number_format = "#,##0.00;[Red]-#,##0.00"
    sheet["N7"] = "；".join(dict.fromkeys(note for note in notes if note)) or "水电资料齐全"
    sheet.merge_cells("N7:AC7")
    sheet["N7"].alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[7].height = max(sheet.row_dimensions[7].height or 15, 56,
                                        16 * ((len(sheet['N7'].value) + 49) // 50))

    image_end = add_images(sheet, images, ocr_results)
    bottom_start = max(last_detail + 3, image_end + 2, 60)
    restore_bottom(sheet, bottom, bottom_start)
    final_bottom = bottom_start + 10
    if sheet.max_row > final_bottom:
        sheet.delete_rows(final_bottom + 1, sheet.max_row - final_bottom)

    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D10"
    sheet.print_area = f"A1:AC{final_bottom}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:9"

    audit = workbook.create_sheet("取数记录")
    header = ["阶段", "月份", "结果", "来源文件", "工作表/位置", "源公式或说明"]
    audit.append(header)
    for row in rent_data["audit"][1:]:
        audit.append((row + [None] * 6)[:6])
    for row in trace:
        audit.append([str(value) if isinstance(value, (Decimal, Path)) else value for value in row])
    for item in details:
        if item.get("source_kind") not in ("普通月度水电", "尾期水电"):
            continue
        audit.append([
            item["source_kind"], item["month"], f"房间{item['room']}：水{item['water']} / 电{item['electric']}",
            Path(item["file"]).name, f"{item['sheet']}!第{item['row']}行", item["content"],
        ])
    for result in ocr_results:
        audit.append(["凭证OCR", "", f"图片{result['number']} / {result['kind']}",
                      Path(app["application_sheet"]).name, f"申请行{app['application_rows']}", result["text"]])
    for cell in audit[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in enumerate([16, 14, 42, 55, 45, 90], 1):
        audit.column_dimensions[get_column_letter(col)].width = width
    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        audit.row_dimensions[row[0].row].height = 34
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    audit.sheet_view.showGridLines = False

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()

    totals = {
        "D6": rent_data["lease_total"], "E6": rent_data["comp_total"],
        "F6": -monthly_values["net_water"], "G6": Decimal("0.00"),
        "H6": rent_data["lease_total"] + rent_data["comp_total"] + monthly_values["net_water"],
        "N14": monthly_values["net_water"],
    }
    for row, key in ((9, "reserve"), (10, "electric_actual"), (11, "water_actual"),
                     (12, "electric_income"), (13, "water_income")):
        totals[f"AC{row}"] = sum(
            (monthly_values.get(month, {}).get(key, Decimal("0.00")) for month in range(1, 9)), Decimal(0)
        )
    patch_formula_cache(output, totals)


def safe_name(value):
    return re.sub(r"[\\/*?:\"<>|]", "_", value)


def generate(args, rules):
    """主流程：租金结果→普通水电→尾期水电→预留→凭证→模板输出。"""
    root = Path(args.root).resolve() if args.root else ROOT
    output = Path(args.output).resolve() if args.output else None
    app = None
    trace = []
    try:
        app_path, history_path, month_paths, app = read_application(root, args)
        exit_month = app["exit_date"].month
        if exit_month > 8:
            raise WaterIssue("当前模板只提供1—8月水电栏位，退出月超过8月时需先扩展模板")
        if output is None:
            filename = safe_name(app["project"] + "_" + app["building"])
            output = root / "水电输出" / f"{filename}_{datetime.now():%Y%m%d_%H%M%S}_含水电对账单.xlsx"
        output.parent.mkdir(parents=True, exist_ok=True)
        if output.exists():
            raise WaterIssue(f"输出已存在：{output}")

        with tempfile.TemporaryDirectory(prefix="water_bill_", dir=output.parent) as temp_dir:
            if args.rent_output:
                rent_output = Path(args.rent_output).resolve()
                if not rent_output.exists():
                    raise WaterIssue(f"指定租金账单不存在：{rent_output}")
                rent_source_note = "复用已生成租金账单，不重新执行租金判断"
            else:
                rent_output = Path(temp_dir) / "租金中间结果.xlsx"
                rent_args = argparse.Namespace(
                    root=str(root), project_name=args.project_name, building_name=args.building_name,
                    building_id=args.building_id, year=args.year, application_row=args.application_row,
                    detail_sheet=args.detail_sheet, claim_type=args.claim_type,
                    claim_details=args.claim_details, output=str(rent_output),
                )
                rent_excel.generate(rent_args, rent_rules)
                rent_source_note = "本次重新执行租金工作流"
            rent_data = read_rent_output(rent_output)
            if rent_data["details"]:
                first = rent_data["details"][0]
                if (rent_excel.name_key(first["project"]) != rent_excel.name_key(app["project"])
                        or rent_excel.name_key(first["building"]) != rent_excel.name_key(app["building"])):
                    raise WaterIssue("指定租金账单与本次项目/楼栋不一致")
            trace.append(["租金基准", "", f"{len(rent_data['details'])}条",
                          rent_output.name, "租金账单/取数记录", rent_source_note])

            monthly_water = []
            zero_notes = []
            for bill_month in range(2, exit_month + 1):
                income_month = bill_month - 1
                if bill_month not in month_paths:
                    raise WaterIssue(f"缺少{bill_month}月月度账单，无法取得{income_month}月水电收入")
                rows, sheet = read_regular_month_water(month_paths[bill_month], bill_month, app, rules)
                monthly_water.extend(rows)
                zero_note = zero_water_note(rows, income_month, f'普通月度水电（取自{bill_month}月账单）')
                if zero_note:
                    zero_notes.append(zero_note)
                trace.append(["普通水电收入", f"{app['exit_date'].year}-{income_month:02d}",
                              f"{len(rows)}条；水{sum((r['water'] for r in rows), Decimal(0))}；电{sum((r['electric'] for r in rows), Decimal(0))}",
                              month_paths[bill_month].name, sheet,
                              f"{bill_month}月账单归属{income_month}月水电收入；仅普通房间明细；不读取延迟及调整页"])

            tail_rows, tail_path, tail_note = read_tail_water(root, app, rules, getattr(args, 'tail_water_input', None))
            monthly_water.extend(tail_rows)
            if not tail_note:
                zero_note = zero_water_note(tail_rows, exit_month, '尾期水电')
                if zero_note:
                    zero_notes.append(zero_note)
            trace.append(["尾期水电", f"{app['exit_date']:%Y-%m}",
                          f"{len(tail_rows)}条；水{sum((r['water'] for r in tail_rows), Decimal(0))}；电{sum((r['electric'] for r in tail_rows), Decimal(0))}",
                          tail_path.name if tail_path else "未找到匹配文件", "", tail_note or "按文件名日期选择最新文件"])

            payments, history_sheet, history_row = read_monthly_paid(history_path, app, exit_month)
            monthly_values = {}
            for month in range(1, exit_month + 1):
                if payments[month]['amount'] <= 0:
                    monthly_values[month] = {'reserve': Decimal('0.00')}
                    trace.append(['水电预留', f"{app['exit_date'].year}-{month:02d}",
                                  f"实付{payments[month]['amount']}；计入0", history_path.name,
                                  f'{history_sheet}!第{history_row}行', '实付≤0，按规则水电预留为0，无需读取月度源预留'])
                    zero_notes.append(f'{month}月水电预留为0（历史实付≤0）')
                    continue
                try:
                    if month not in month_paths:
                        raise ReserveUnavailable(f'缺少{month}月月度账单')
                    source_reserve, sheet, cell, formula = read_source_reserve(month_paths[month], month, app, rules)
                except (ReserveUnavailable, FileNotFoundError) as exc:
                    note = f"{month}月历史实付{payments[month]['amount']}元，未取得源水电预留，暂按0；{exc}；待核对"
                    monthly_values[month] = {'reserve': Decimal('0.00')}
                    zero_notes.append(note)
                    trace.append(['水电预留缺失', f"{app['exit_date'].year}-{month:02d}", note,
                                  str(month_paths.get(month, '')), '', '暂按0，非已核实零额'])
                    continue
                recorded = rules.reserve_to_record(payments[month]["amount"], source_reserve)
                monthly_values[month] = {"reserve": recorded}
                trace.append(["水电预留", f"{app['exit_date'].year}-{month:02d}",
                              f"实付{payments[month]['amount']}；源预留{source_reserve}；计入{recorded}",
                              month_paths[month].name, f"{sheet}!{cell}",
                              f"付款来源{history_path.name}/{history_sheet}!第{history_row}行；{formula}"])

            filled_water, filled_electric, source_cells, images = read_application_actuals_and_images(app_path, app, rules)
            if not images:
                from boyu_agent.vouchers import fallback_vouchers
                images, voucher_notes = fallback_vouchers(root, app)
                zero_notes.extend(voucher_notes)
                for item in images:
                    trace.append(['外部划扣凭证', '', item.get('source', ''), item.get('source', ''), '', '申请行无图片，项目及楼栋匹配'])
            ocr_results, ocr_note = ocr_images(images, skip=args.skip_ocr)
            actual_water, actual_electric, voucher_water, compare_notes = compare_actuals(
                filled_water, filled_electric, ocr_results, app["exit_date"], rules
            )
            notes = rent_data.get('notes', []) + [tail_note, ocr_note] + zero_notes + compare_notes
            trace.append(["划扣填写", "1—4月", f"水{filled_water}；电{filled_electric}",
                          app_path.name, f"{app['application_sheet']}!第{app['application_rows'][-1]}行", str(source_cells)])
            trace.append(["凭证核验", "", f"图片{len(images)}张；OCR水费{voucher_water}",
                          app_path.name, f"{app['application_sheet']}!第{app['application_rows'][-1]}行", "；".join(notes)])

            income_by_month = defaultdict(lambda: {"water": Decimal(0), "electric": Decimal(0)})
            for row in monthly_water:
                month = int(row["month"][-2:])
                income_by_month[month]["water"] += row["water"]
                income_by_month[month]["electric"] += row["electric"]
            for month in range(1, exit_month + 1):
                monthly_values[month].update({
                    "water_actual": actual_water.get(month, Decimal("0.00")),
                    "electric_actual": actual_electric.get(month, Decimal("0.00")),
                    "water_income": income_by_month[month]["water"],
                    "electric_income": income_by_month[month]["electric"],
                })
            reserve_total = sum((monthly_values[m]["reserve"] for m in range(1, exit_month + 1)), Decimal(0))
            actual_total = sum((monthly_values[m]["water_actual"] + monthly_values[m]["electric_actual"]
                                for m in range(1, exit_month + 1)), Decimal(0))
            income_total = sum((monthly_values[m]["water_income"] + monthly_values[m]["electric_income"]
                                for m in range(1, exit_month + 1)), Decimal(0))
            monthly_values["net_water"] = reserve_total - actual_total + income_total

            details = merge_details(rent_data["details"], monthly_water)
            trace.insert(0, ["工作流", "", rules.WORKFLOW_VERSION, "", "", "租金工作流版本=" + rent_rules.WORKFLOW_VERSION + "；水电按普通明细取数"])
            trace.append(["水电结算", "", f"预留{reserve_total}-实际支出{actual_total}+收入{income_total}={monthly_values['net_water']}",
                          "", "模板N14/F6", "F6=-N14"])
            write_output(
                root / "输出指引" / "模板文件_尾期对账单输出.xlsx", output, app,
                rent_data, details, monthly_values, notes, images, ocr_results, trace, rules,
            )

        return {
            "version": rules.WORKFLOW_VERSION,
            "output": str(output),
            "status": "含水电对账单已生成",
            "rent_rows": len(rent_data["details"]),
            "water_rows": len(monthly_water),
            "tail_water_rows": len(tail_rows),
            "tail_note": tail_note,
            "reserve_total": str(reserve_total),
            "actual_water_electric": str(actual_total),
            "water_electric_income": str(income_total),
            "net_water": str(monthly_values["net_water"]),
            "final_building_share": str(rent_data["lease_total"] + rent_data["comp_total"] + monthly_values["net_water"]),
            "notes": list(dict.fromkeys(note for note in notes if note)),
        }
    except Exception as exc:
        target = (output.with_suffix(".诊断.json") if output else root / "水电输出" / f"取数异常_{datetime.now():%Y%m%d_%H%M%S}.json")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps({"error": str(exc), "application": app, "trace": trace},
                                     ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        raise WaterIssue(f"{exc}\n诊断文件：{target}") from exc
