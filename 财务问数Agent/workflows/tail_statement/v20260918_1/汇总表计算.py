"""最终汇总：间次计费、完整备注、历史单据、数据来源。"""
from pathlib import Path
from datetime import datetime
from decimal import Decimal
import re
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
import 租金Excel接入 as rent
import 租金板块工作流 as rent_rules
from 水电Excel接入 import patch_formula_cache, TAIL_FOLDER

WORKFLOW_VERSION = '2026-09-18.1'


class SummaryIssue(ValueError):
    """需要核实的判断问题，不自动猜测。"""


def format_summary_header(sheet):
    sheet.row_dimensions[5].height=62
    sheet.row_dimensions[6].height=48
    sheet['C5']='计费在租房间数\n（累计间次）'
    for row in (5,6):
        for col in range(1,12):
            cell=sheet.cell(row,col)
            if cell.__class__.__name__=='MergedCell': continue
            cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    sheet['A6'].font=Font(name='微软雅黑',size=10)
    sheet['N7'].font=Font(name='微软雅黑',size=11)


def latest_named(names, pattern):
    matches = [(datetime.strptime(m[1], '%Y%m%d'), name) for name in names
               if (m := re.fullmatch(pattern, name))]
    if not matches:
        raise SummaryIssue('未找到日期命名的最新数据：'+str(names))
    newest = max(day for day, _ in matches)
    hits = [name for day, name in matches if day == newest]
    if len(hits) != 1:
        raise SummaryIssue('同日最新数据不唯一：'+str(hits))
    return hits[0]


def fee_months(audit, exit_month, exit_type):
    selected = {}
    for row in audit:
        if len(row) < 6 or row[0] != '选择明细':
            continue
        month = int(str(row[1])[-2:])
        if not re.match(r'[1-9]\d*条', str(row[2])):
            continue
        if (month < exit_month and row[5] == '历史实付为0') or (month == exit_month and exit_type == 'compensation'):
            selected[month] = str(row[5])
    return selected


def read_fee_counts(month_paths, app, selected):
    output = []
    for month, reason in sorted(selected.items()):
        if month not in month_paths:
            raise SummaryIssue(f'{month}月需计费，但月度文件缺失，需核实在租房间数')
        book = rent.SourceBook(month_paths[month])
        try:
            sheet = rent.SUMMARY_NAMES.get(month, '汇总表')
            headers, rows = rent.table(book, sheet, ['id'])
            id_i = rent.index(headers, rent.ALIASES['id'])
            count_i = rent.index(headers, ['0101在租资源数', '在租房间数', '在租资源数'])
            hits = [(n, r) for n, r, _ in rows if rent.same_id(rent.cell(r, id_i), app['building_id'])]
            if len(hits) != 1:
                raise SummaryIssue(f'{month}月计费房间数匹配{len(hits)}条，请核实')
            row_no, row = hits[0]
            count = rent_rules.amount(rent.cell(row, count_i))
            if count < 0 or count != count.to_integral_value():
                raise SummaryIssue(f'{month}月在租房间数无效：{count}')
            output.append(dict(month=month, count=int(count), fee=count*100, reason=reason,
                               file=str(book.path.resolve()), location=f'{sheet}!{get_column_letter(count_i+1)}{row_no}'))
        finally:
            book.close()
    return output


def document_months(headers, row):
    mapping = {}
    for index, header in enumerate(headers):
        if '单据' not in str(header) or rent_rules.is_draft_header(header):
            continue
        value = str(rent.cell(row, index) or '').strip()
        if not value or value in ('0', '0.0', '0.00'):
            continue
        match = re.match(r'\s*(\d{1,2})月', str(header))
        month = int(match[1]) if match else '尾期/未标明月份'
        codes = re.findall(r'NP[A-Z0-9]+', value.upper())
        if not codes:
            raise SummaryIssue(f'付款申请列{header}单据号无法识别：{value}')
        for code in codes:
            if code in mapping and mapping[code][0] != month:
                raise SummaryIssue(f'同一单据{code}对应多个月份，请核实')
            mapping[code] = (month, get_column_letter(index+1))
    return mapping


def read_documents(root, history_path, app):
    hb = rent.SourceBook(history_path)
    try:
        history_sheet = rent_rules.select_history_sheet(list(hb.sheets))
        headers, rows = rent.table(hb, history_sheet, ['id'])
        id_i = rent.index(headers, rent.ALIASES['id'])
        hits = [(n, r) for n, r, _ in rows if rent.same_id(rent.cell(r,id_i), app['building_id'])]
        if len(hits) != 1:
            raise SummaryIssue('付款申请表楼栋无法唯一匹配')
        history_row, row = hits[0]
        mapping = document_months(headers, row)
    finally:
        hb.close()
    database = root/rent.DATABASE_FOLDER
    paths = {p.name:p for p in database.glob('*付款单据.xlsx') if not p.name.startswith('~$')}
    selected = latest_named(list(paths), r'.*截止(\d{8})付款单据\.xlsx')
    path = paths[selected]
    book = rent.SourceBook(path)
    try:
        sheet = latest_named(list(book.sheets), r'(\d{8})最新')
        rows = book.rows(sheet)
        headers = rows[0][1]
        cols = {key:rent.index(headers,[label]) for key,label in {
            'code':'报销编码','subject':'关联事项','amount':'申请报销金额','time':'事项申请时间'}.items()}
        found = {}
        for row_no, row, _ in rows[1:]:
            code = str(rent.cell(row,cols['code']) or '').strip().upper()
            if code not in mapping:
                continue
            record = {key:rent.cell(row,index) for key,index in cols.items()}
            value = record['time']
            if isinstance(value,(Decimal,int,float)):
                record['time'] = from_excel(float(value))
            elif value:
                record['time'] = datetime.fromisoformat(str(value))
            record['amount'] = rent_rules.amount(record['amount'])
            record.update(month=mapping[code][0], row=row_no, file=str(path.resolve()), sheet=sheet,
                          month_source=f'{history_sheet}!{mapping[code][1]}{history_row}')
            if code in found:
                old = found[code]
                if old['amount'] != record['amount'] or old['subject'] != record['subject']:
                    raise SummaryIssue(f'单据{code}重复记录的金额/关联事项不一致，请核实')
                # 同号同金额同事项只保留首条，源时间随首条保留。
                continue
            found[code] = record
        missing = [f'历史单据{code}（{month}月）未在{sheet}取得' for code,(month,_) in mapping.items() if code not in found]
        records = sorted(found.values(),key=lambda r:(r['month'] if isinstance(r['month'],int) else 99, r['code']))
        return records, missing, path, sheet
    finally:
        book.close()


def read_contract_mode(root, app):
    paths=[p for p in (root/rent.DATABASE_FOLDER).glob('*万村代管*付款申请*.xlsx') if not p.name.startswith('~$')]
    if len(paths)!=1:
        raise SummaryIssue(f'合同模式底表匹配{len(paths)}个文件，请核实')
    book=rent.SourceBook(paths[0])
    sheet='(底表勿动)1-7月付款合同与计划'
    try:
        headers, rows=rent.table(book,sheet,['id'])
        id_i=rent.index(headers,rent.ALIASES['id'])
        type_i=rent.index(headers,['合同类型'])
        hits=[(n,str(rent.cell(row,type_i) or '').strip()) for n,row,_ in rows
              if rent.same_id(rent.cell(row,id_i),app['building_id'])]
        if not hits:
            raise SummaryIssue('合同模式底表未取得该楼栋ID：'+app['building_id'])
        raw=[value for _,value in hits if value]
        if any('分成' in value for value in raw):
            mode='分成'
        elif raw and all(('代管' in value or '声明函' in value) for value in raw):
            mode='代管'
        else:
            raise SummaryIssue(f'合同类型无法归类：{raw or ["空白"]}')
        return mode, paths[0], sheet, hits
    finally:
        book.close()


def build_high_risks(app, audit, extra, contract_mode):
    risks=[]
    if contract_mode=='分成':
        risks.append('该楼栋合同模式为分成，需单独核对')
    for row in audit:
        if not row:
            continue
        kind=str(row[0] or '')
        if kind=='高风险-历史付款差异':
            risks.append(str(row[5]))
        elif kind=='高风险-退出月租金缺失':
            risks.append(str(row[2]))
        elif kind=='租金缺失备注' and str(row[1])!=f'{app["exit_date"].year}-{app["exit_date"].month:02d}':
            risks.append(str(row[2])+'，待核对')
    keywords=('不一致','未取得对应凭证','未取得尾期水电','未核验','待核对')
    for item in extra:
        text=str(item or '').strip()
        if text and any(word in text for word in keywords):
            risks.append(text if text.endswith('待核对') else text+'，待核对')
    return list(dict.fromkeys(risks))


def build_notes(app, audit, values, fee_rows, extra, contract_mode='代管'):
    selected = [f'{r[1]}：{r[5]}' for r in audit if len(r)>5 and r[0]=='选择明细']
    excluded = [f'{r[1]}：{r[5]}' for r in audit if len(r)>5 and r[0]=='后续月排除']
    paid = [f'{r[1]}已支付' for r in audit if len(r)>2 and r[0]=='历史付款' and Decimal(str(r[2]))!=0]
    delayed = [str(r[1]) for r in audit if r and r[0]=='延迟汇总核对']
    missing = [str(r[2]) for r in audit if len(r)>2 and r[0]=='租金缺失备注']
    follow = []
    for r in audit:
        if len(r)>5 and r[0]=='选择明细' and '后续月' in str(r[5]):
            follow.append(f'{r[1]}已取（租赁收入＋赔偿金≤0）')
    water = [f'{r[1]}{r[0]}：{r[2]}' for r in audit if len(r)>2 and r[0] in ('普通水电收入','尾期水电') and re.match(r'\d+条',str(r[2]))]
    actual = []
    for month in range(1,app['exit_date'].month+1):
        col = 14+(month-1)*2
        if col>28:
            raise SummaryIssue('当前模板水电区仅支持1—8月，需扩展模板后再处理该退出月份')
        electric = values.cell(10,col).value or 0
        water_paid = values.cell(11,col).value or 0
        actual.append(f'{month}月水{water_paid}、电{electric}')
    fees = '；'.join(f"{r['month']}月{r['count']}间次/{r['fee']}元" for r in fee_rows) or '无符合条件月份，0间次/0元'
    charged = {r['month']: r['fee'] for r in fee_rows}
    zero_fee_months = [str(month)+'月' for month in range(1, app['exit_date'].month+1)
                       if charged.get(month, 0) == 0]
    if zero_fee_months:
        fees += '；管理费为0的月份：'+'、'.join(zero_fee_months)+'（本次未计管理费）'
    return [
        '一、合同模式：'+contract_mode+'；退出：'+('赔偿' if app['exit_type']=='compensation' else '免租')+f"；{app['exit_date']:%Y年%m月}。",
        '二、月度租金：'+('；'.join(selected) or '未提取')+'。未取：'+('；'.join(paid+excluded) or '无')+'。延迟明细月份：'+('、'.join(dict.fromkeys(delayed)) or '无')+'。计费：'+fees+'。',
        '三、后续月及缺失：'+('；'.join(follow+excluded+missing) or '无后续月缺失/零额项目')+'。',
        '四、水电取数：'+'；'.join(water)+'。',
        '五、实际水电支出：'+'；'.join(actual)+'。1—4月按申请填写及凭证核验；后续月份按既有水电规则，未取得按0。',
        '六、其他异常：'+('；'.join(dict.fromkeys(str(x) for x in extra if x)) or '无')+'。',
    ]


def finish(root, args, water_output, output, water_result):
    app_path, history_path, month_paths = rent.discover(root)
    ab = rent.SourceBook(app_path)
    try:
        app = rent.resolve_application(ab,args,rent_rules)
    finally:
        ab.close()
    workbook = load_workbook(water_output,keep_links=False)
    cached = load_workbook(water_output,data_only=True,keep_links=False)
    try:
        sheet = workbook.worksheets[0]
        values = cached.worksheets[0]
        audit_sheet = workbook['取数记录']
        audit = [[c.value for c in row] for row in audit_sheet]
        selected = fee_months(audit,app['exit_date'].month,app['exit_type'])
        fee_rows = read_fee_counts(month_paths,app,selected)
        documents, doc_notes, document_path, document_sheet = read_documents(root,history_path,app)
        contract_mode, contract_path, contract_sheet, contract_rows = read_contract_mode(root,app)
        extra=water_result['notes']+doc_notes
        risks=build_high_risks(app,audit,extra,contract_mode)
        notes = build_notes(app,audit,values,fee_rows,extra,contract_mode)
        count = sum(r['count'] for r in fee_rows)
        fee = Decimal(count)*100
        final = Decimal(water_result['final_building_share'])-fee
        sheet['A1']=f"（机处理）{app['project']} {app['building']} 对账单最终版"
        sheet['C5']='计费在租房间数\n（各月累计间次）'
        sheet['C6']=count
        sheet['G6']='=C6*100'
        sheet['H6']='=D6+E6-F6-G6'
        sheet['K6']='（机处理）对账单最终版'
        risk_text=['高风险提示：']+([f'高风险提示{i}：{value}' for i,value in enumerate(risks,1)] if risks else ['无'])
        sheet['N7']='\n'.join(risk_text+['','备注：']+notes)
        sheet['N7'].alignment=Alignment(wrap_text=True,vertical='top')
        sheet.row_dimensions[7].height=min(409, max(140, 15*sum((len(n)+70)//71 for n in notes)))
        format_summary_header(sheet)
        # 右侧历史付款表与数据来源安排在已有明细、凭证下方，避免遮挡。
        start = sheet.max_row+3
        for address in list(sheet.merged_cells.ranges):
            if address.min_col>=13 and address.min_row>=31:
                sheet.unmerge_cells(str(address))
        blue='1F4E78'
        def title(row,text):
            sheet.merge_cells(start_row=row,start_column=13,end_row=row,end_column=29)
            cell=sheet.cell(row,13,text); cell.font=Font(name='微软雅黑',bold=True,color='FFFFFF',size=11)
            cell.fill=PatternFill('solid',fgColor=blue)
            cell.alignment=Alignment(vertical='center'); sheet.row_dimensions[row].height=24
        def line(row, entries, height=36, header=False):
            for left,right,value in entries:
                if right>left: sheet.merge_cells(start_row=row,start_column=left,end_row=row,end_column=right)
                cell=sheet.cell(row,left,value)
                cell.font=Font(name='微软雅黑',size=10,bold=header)
                cell.alignment=Alignment(wrap_text=True,vertical='center')
                cell.fill=PatternFill('solid',fgColor='DCE6F1' if header else 'FFFFFF')
                cell.border=Border(bottom=Side(style='thin',color='D9E2F3'))
            sheet.row_dimensions[row].height=height
        title(start,'历史付款明细（不按付款状态筛选）')
        line(start+1,[(13,13,'月份'),(14,17,'单据号'),(18,23,'关联事项'),(24,26,'申请金额'),(27,29,'事项申请时间')],26,True)
        pos=start+2
        for record in documents:
            label=f"{record['month']}月" if isinstance(record['month'],int) else record['month']
            line(pos,[(13,13,label),(14,17,record['code']),(18,23,record['subject']),(24,26,float(record['amount'])),(27,29,record['time'])],42)
            sheet.cell(pos,24).number_format='#,##0.00'
            sheet.cell(pos,27).number_format='yyyy-mm-dd hh:mm:ss'
            audit_sheet.append(['历史单据',label,record['code'],document_path.name,f"{document_sheet}!行{record['row']}",f"月份来源：{history_path.name}/{record['month_source']}"])
            pos+=1
        if not documents:
            line(pos,[(13,29,'未取得可匹配的历史单据；详见备注')]); pos+=1
        pos+=1
        title(pos,'我方应收费用计算明细'); pos+=1
        line(pos,[(13,14,'月份'),(15,17,'在租间次'),(18,20,'单价'),(21,23,'应收费用'),(24,29,'取数位置')],26,True);pos+=1
        for row in fee_rows:
            line(pos,[(13,14,f"{row['month']}月"),(15,17,row['count']),(18,20,100),(21,23,float(row['fee'])),(24,29,row['location'])]);pos+=1
            audit_sheet.append(['最终汇总计费',f"{row['month']}月",f"{row['count']}间次×100={row['fee']}",Path(row['file']).name,row['location'],row['reason']])
        if not fee_rows:
            line(pos,[(13,29,'无符合条件的计费月份，0间次，0元')]);pos+=1
        source_names={str(r[3]) for r in audit if len(r)>3 and r[3]}
        sources={app_path.resolve(),history_path.resolve(),document_path.resolve(),contract_path.resolve(),(root/'输出指引/模板文件_尾期对账单输出.xlsx').resolve()}
        sources.update(p.resolve() for month,p in month_paths.items() if app['exit_type']=='compensation' or month<=app['exit_date'].month)
        sources.add((root/'输出指引/对账单出账判断流程.xmind').resolve())
        sources.update(Path(r['file']).resolve() for r in fee_rows)
        for folder in (rent.MONTHLY_FOLDER,rent.DATABASE_FOLDER,TAIL_FOLDER,'输出指引'):
            sources.update(p.resolve() for p in (root/folder).rglob('*.xlsx') if p.name in source_names)
        pos+=1; title(pos,'数据来源'); pos+=1
        line(pos,[(13,14,'序号'),(15,29,'实际使用文件（相对工作根目录）')],26,True);pos+=1
        for i,path in enumerate(sorted(sources,key=str),1):
            line(pos,[(13,14,i),(15,29,str(path.relative_to(root)))],32);pos+=1
        sheet.print_area=f'A1:AC{pos}'
        # 将全部备注另存为逐项记录，方便查看较长的缺失说明。
        for note in notes:
            audit_sheet.append(['最终备注','',note,'','',''])
        for risk in risks:
            audit_sheet.append(['高风险提示','',risk,contract_path.name if '合同模式' in risk else '','',''])
        audit_sheet.append(['合同模式','',contract_mode,contract_path.name,contract_sheet,
                            '匹配行：'+','.join(str(row) for row,_ in contract_rows)])
        audit_sheet.append(['最终汇总版本','',WORKFLOW_VERSION,'','',''])
        audit_sheet.auto_filter.ref=audit_sheet.dimensions
        cache_values={c.coordinate:values[c.coordinate].value for row in sheet for c in row
                      if c.data_type=='f' and values[c.coordinate].value is not None}
        cache_values.update(G6=fee,H6=final)
        output.parent.mkdir(parents=True,exist_ok=True)
        if output.exists(): raise SummaryIssue(f'输出已存在：{output}')
        workbook.save(output)
        patch_formula_cache(output,cache_values)
        return dict(version=WORKFLOW_VERSION,output=str(output),room_count=count,management_fee=str(fee),
                    final_building_share=str(final),history_documents=len(documents),contract_mode=contract_mode,
                    high_risks=risks,notes=notes)
    finally:
        cached.close();workbook.close()
